"""The .xlsx export of one stored backtest, laid out like the user's manual sheet:
title row, then DATE | CALL/PUT | 15 minute candel size | PROFIT | LOSS | EOD, one row per trading day.

EOD follows the user's sheet: the size of the day's trade result in % (a loss row shows the loss size,
a profit row the profit size) - the PROFIT/LOSS columns carry the sign.
"""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

TITLES = {"NIFTY": "NIFTY 50", "BANKNIFTY": "BANK NIFTY"}
HEADERS = ["DATE", "CALL/PUT", "15 minute candel size", "PROFIT", "LOSS", "EOD"]
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _row(r: dict) -> list:
    pnl = r.get("pnl_points")
    side = {"CALL": "CALL", "PUT": "PUT", "BOTH": "UNRESOLVED"}.get(r.get("side") or "", "")
    if r["exit_reason"] == "NO_ENTRY":
        side = "NO ENTRY"
    elif r["exit_reason"].startswith("DATA") and pnl is None:
        side = "DATA INCOMPLETE"
    size = None if r.get("first_range_pct") is None else r["first_range_pct"] / 100.0
    eod = None if r.get("pnl_pct") is None else abs(r["pnl_pct"]) / 100.0
    return [
        dt.date.fromisoformat(r["date"]), side, size,
        "PROFIT" if pnl is not None and pnl > 0 else None,
        "LOSS" if pnl is not None and pnl < 0 else None,
        eod,
    ]


def build_xlsx(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:F2")
    ws["A1"] = TITLES.get(result["symbol"], result["symbol"])
    ws["A1"].font = Font(bold=True, size=20)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEF":
        for row in (1, 2):
            ws[f"{col}{row}"].border = BOX
    for c, text in zip(ws[3], HEADERS):
        c.value = text
        c.font = Font(bold=True)
        c.border = BOX
    for n, r in enumerate(result["rows"], start=4):
        cells = ws[n]
        for c, v in zip(cells, _row(r)):
            c.value = v
        cells[0].number_format = "d-mmm-yy"
        cells[2].number_format = "0.00%"
        cells[5].number_format = "0.00%"
        for c in cells:
            c.border = BOX
    for letter, width in zip("ABCDEF", (12, 11, 24, 10, 10, 10)):
        ws.column_dimensions[letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
