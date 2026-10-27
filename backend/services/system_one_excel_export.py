"""Excel export for a finished System One backtest. Writes exactly the
stored result (same rows, same column order as the UI table); never
recomputes anything."""
from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.services.system_one_backtest_service import INTRADAY_TIMEFRAMES, PNL_NOTE, StoredSystemOneBacktest

HEADER_FILL = PatternFill("solid", fgColor="FFD966")
HEADER_FONT = Font(bold=True)
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREEN_FONT = Font(color="006100")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9") for _ in range(4)))

DATE_FORMAT = "dd-mmm-yyyy"
DATETIME_FORMAT = "dd-mmm-yyyy hh:mm"
MONEY_FORMAT = "#,##0.00"
PNL_FORMAT = "#,##0.00;-#,##0.00"

# Same order as the table in the UI (frontend SystemOneBacktestPage.tsx).
TRADE_HEADERS = [
    "Trade #", "Signal Date", "Entry Date", "Entry Price", "Initial Stop (MA at entry)", "Risk Per Share",
    "Quantity", "Planned Risk (Rs)", "Exit Date", "Exit Price", "Exit Reason", "Holding Weeks",
    "P&L (Rs)", "P&L % of Entry", "R-Multiple", "Unrealised P&L (open)", "Data-Quality Flag",
]
_PNL_COL = TRADE_HEADERS.index("P&L (Rs)") + 1


def file_name(stored: StoredSystemOneBacktest) -> str:
    return f"system_one_{stored.symbol}_{stored.timeframe}_{stored.from_date:%Y%m%d}_{stored.to_date:%Y%m%d}.xlsx"


def _autosize(ws, min_width: int = 12, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, min_width), max_width)


def build_excel(stored: StoredSystemOneBacktest) -> bytes:
    wb = openpyxl.Workbook()

    # ---- Sheet 1: Trades — one row per trade, header on row 1 ---------------
    ws = wb.active
    ws.title = "Trades"
    # Intraday candles keep their time of day; daily/weekly candles are plain dates.
    intraday = stored.timeframe in INTRADAY_TIMEFRAMES
    date_format = DATETIME_FORMAT if intraday else DATE_FORMAT

    def _when(ts):
        return None if ts is None else (ts.to_pydatetime() if intraday else ts.date())

    ws.append(TRADE_HEADERS)
    for c in range(1, len(TRADE_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
        cell.alignment = Alignment(horizontal="center")

    for t in stored.result.trades:
        # Open trades leave the exit cells empty - never 0; their unrealised P&L has its own column.
        ws.append([
            t.trade_number, _when(t.signal_date), _when(t.entry_date), t.entry_price, t.stop_loss, t.risk_per_share,
            t.quantity, t.risk_per_trade, _when(t.exit_date), t.exit_price, t.exit_reason, t.holding_periods,
            t.pnl_amount, t.pnl_pct, t.r_multiple, t.unrealized_pnl, t.data_flag,
        ])
        row = ws.max_row
        for col, fmt in ((2, date_format), (3, date_format), (4, MONEY_FORMAT), (5, MONEY_FORMAT), (6, MONEY_FORMAT),
                         (7, "0"), (8, MONEY_FORMAT), (9, date_format), (10, MONEY_FORMAT), (13, PNL_FORMAT),
                         (14, "0.00"), (15, "0.00"), (16, PNL_FORMAT)):
            ws.cell(row=row, column=col).number_format = fmt
        for col in range(1, len(TRADE_HEADERS) + 1):
            ws.cell(row=row, column=col).border = THIN_BORDER
        pnl_cell = ws.cell(row=row, column=_PNL_COL)
        if t.pnl_amount is not None and t.pnl_amount > 0:
            pnl_cell.fill, pnl_cell.font = GREEN_FILL, GREEN_FONT
        elif t.pnl_amount is not None and t.pnl_amount < 0:
            pnl_cell.fill, pnl_cell.font = RED_FILL, RED_FONT
    ws.freeze_panes = "A2"
    _autosize(ws)

    # ---- Sheet 2: Summary, settings and skipped signals ---------------------
    s = stored.result.summary()
    summary = wb.create_sheet("Summary")
    summary.append(["Field", "Value"])
    rows = [
        ("Strategy", "System One"),
        ("Symbol", stored.symbol),
        ("Trading Symbol", stored.trading_symbol),
        ("Timeframe", stored.timeframe),
        ("From Date", stored.from_date.date()),
        ("To Date", stored.to_date.date()),
        ("Data Available From", stored.data_available_from.date()),
        ("Data Available To", stored.data_available_to.date()),
        ("Reference Capital", stored.config.reference_capital),
        ("Risk Per Trade %", stored.config.risk_per_trade_pct),
        ("Risk Per Trade", stored.config.risk_per_trade),
        ("Total Trades", s.total_trades),
        ("Closed Trades", s.closed_trades),
        ("Open Trades", s.open_trades),
        ("Winning Trades", s.winning_trades),
        ("Losing Trades", s.losing_trades),
        ("Net Realised P&L (closed trades only)", s.net_realized_pnl),
        ("Win Rate % (closed trades)", s.win_rate),
        ("Average Win", s.average_win),
        ("Average Loss", s.average_loss),
        ("Profit Factor", s.profit_factor if s.profit_factor is not None else "N/A"),
        ("Best Trade", s.best_trade),
        ("Worst Trade", s.worst_trade),
        ("Max Drawdown (Rs, realised)", s.max_drawdown),
        ("Longest Losing Streak", s.longest_losing_streak),
        ("Unrealised P&L (open trades, not in totals)", s.unrealized_pnl),
        ("Note", PNL_NOTE),
    ]
    for label, value in rows:
        summary.append([label, value])
        cell = summary.cell(row=summary.max_row, column=2)
        if hasattr(value, "year"):
            cell.number_format = DATE_FORMAT
        elif isinstance(value, float):
            cell.number_format = PNL_FORMAT
    for w in stored.result.warnings:
        summary.append(["Warning", w])
    if stored.result.skipped:
        summary.append([])
        summary.append(["Skipped Signal Date", "Reason"])
        for sk in stored.result.skipped:
            summary.append([_when(sk.signal_date), sk.reason])
            summary.cell(row=summary.max_row, column=1).number_format = date_format
    _autosize(summary, max_width=90)

    # ---- Sheet 3: year-wise (by exit year) and equity curve ------------------
    years = wb.create_sheet("Year-wise")
    years.append(["Year (exit)", "Trades", "Wins", "Losses", "Gross Profit", "Gross Loss", "Net P&L"])
    for y in stored.result.yearly():
        years.append([y.year, y.trades, y.wins, y.losses, y.gross_profit, y.gross_loss, y.net_pnl])
        for col in (5, 6, 7):
            years.cell(row=years.max_row, column=col).number_format = PNL_FORMAT
    _autosize(years)

    curve = wb.create_sheet("Equity Curve")
    curve.append(["Exit Date", "Trade #", "Cumulative Realised P&L"])
    for p in stored.result.equity_curve():
        curve.append([_when(p.date), p.trade_number, p.cumulative_pnl])
        curve.cell(row=curve.max_row, column=1).number_format = date_format
        curve.cell(row=curve.max_row, column=3).number_format = PNL_FORMAT
    _autosize(curve)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
