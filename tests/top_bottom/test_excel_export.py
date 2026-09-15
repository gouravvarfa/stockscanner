import datetime as dt

import openpyxl

from backend.strategies.top_bottom.backtest import run_top_bottom_backtest
from backend.strategies.top_bottom.models import BacktestResult, DataCoverage, PricePoint
from backend.strategies.top_bottom.statistics import build_equity_curve, compute_statistics
from backend.services.top_bottom_excel_export import build_excel


def _dates(n, start="2026-02-01"):
    base = dt.datetime.strptime(start, "%Y-%m-%d")
    return [base + dt.timedelta(days=i) for i in range(n)]


def _make_result(lot_size: int) -> BacktestResult:
    closes = [10, 40, 30, 40, 70, 100, 90, 110, 90, 80]
    dates = _dates(len(closes))
    opens = closes[:1] + closes[:-1]
    signals, trades = run_top_bottom_backtest("BANKNIFTY", "BANKNIFTY29SEP26FUT", "2026-09-29", "INDEX", lot_size, dates, closes, opens)
    stats = compute_statistics(signals, trades)
    equity = build_equity_curve(trades)
    coverage = DataCoverage(dates[0], dates[-1], dates[0], dates[-1], True)
    return BacktestResult(
        backtest_id="test-id", symbol="BANKNIFTY", trading_symbol="BANKNIFTY29SEP26FUT",
        expiry="2026-09-29", instrument_class="INDEX", exch_seg="NFO", timeframe="1D",
        coverage=coverage,
        price_series=[PricePoint(d, c) for d, c in zip(dates, closes)],
        system_point_line=[], stop_loss_line=[],
        signals=signals, trades=trades, equity_curve=equity,
        statistics=stats, starting_capital=100.0, created_at=dt.datetime.now(),
    )


def test_pnl_amount_uses_the_real_lot_size_not_a_guess():
    result = _make_result(lot_size=30)
    trade = result.trades[0]
    assert trade.pnl_amount == trade.pnl_points * 30


def test_pnl_amount_is_none_when_lot_size_unknown_not_fabricated_zero():
    result = _make_result(lot_size=0)
    trade = result.trades[0]
    assert trade.pnl_amount is None


def test_excel_trade_results_sheet_has_the_simple_columns_and_totals():
    result = _make_result(lot_size=30)
    content = build_excel(result)

    wb = openpyxl.load_workbook(__import__("io").BytesIO(content))
    ws = wb["Trade Results"]

    header = [ws.cell(row=2, column=c).value for c in range(1, 7)]
    assert header == ["Date", "Buy/Sell", "Entry", "Exit", "Point", "Profit / Loss"]

    # Totals block (Total Trades and friends) is visible on this same sheet.
    labels = {ws.cell(row=r, column=8).value for r in range(2, 8)}
    assert "Total Trades" in labels
    assert "Net P&L (Rs.)" in labels or any(ws.cell(row=r, column=10).value == "Net P&L (Rs.)" for r in range(2, 8))


def test_excel_has_trade_detail_and_summary_sheets_too():
    result = _make_result(lot_size=30)
    content = build_excel(result)
    wb = openpyxl.load_workbook(__import__("io").BytesIO(content))
    assert set(wb.sheetnames) >= {"Trade Results", "Trade Detail", "Backtest Summary", "Equity Curve", "Signal Log"}
