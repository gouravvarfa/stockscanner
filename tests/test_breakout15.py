import datetime as dt

from backend.strategies.breakout15 import engine as E

D = dt.date(2025, 1, 7)


def t(h, m):
    return dt.datetime.combine(D, dt.time(h, m))


def starts():
    out, x = [], 9 * 60 + 15
    while x < 15 * 60 + 30:
        out.append((x // 60, x % 60))
        x += 15
    return out


def day(first=(100, 110, 90, 105), body=None, skip=(), until=(15, 15)):
    """Full session of 15-minute candles (bar START times). Unlisted candles sit flat at 100."""
    body, bars = body or {}, []
    for hm in starts():
        if hm in skip or hm > until:
            continue
        o, h, l, c = first if hm == (9, 15) else body.get(hm, (100, 100, 100, 100))
        bars.append((t(*hm), o, h, l, c))
    return bars


def test_session_has_25_candles():
    assert len(starts()) == 25 == E.EXPECTED_BARS


def test_first_candle_stats_and_no_entry_day_is_kept():
    r = E.run_day(D, day())
    assert (r.first_open, r.first_high, r.first_low, r.first_close) == (100, 110, 90, 105)
    assert r.first_range_pct == 20.0 and r.first_direction == "Green"
    assert r.exit_reason == E.NO_ENTRY and r.entry_price is None and r.pnl_points is None


def test_red_and_doji_direction():
    assert E.run_day(D, day(first=(100, 110, 90, 95))).first_direction == "Red"
    assert E.run_day(D, day(first=(100, 110, 90, 100))).first_direction == "Doji"


def test_breakout_enters_at_first_high_and_exits_at_market_close():
    r = E.run_day(D, day(body={(10, 0): (108, 111, 108, 110), (15, 15): (115, 116, 114, 120)}))
    assert r.breakout_time == t(10, 0) and r.entry_price == 110
    assert r.exit_reason == E.MARKET_CLOSE and r.exit_price == 120 and r.pnl_points == 10
    assert round(r.pnl_pct, 4) == round(10 / 110 * 100, 4)


def test_opening_gap_enters_at_candle_open():
    r = E.run_day(D, day(body={(9, 30): (112, 113, 111, 112)}))
    assert r.side == "CALL" and r.entry_price == 112 and "gap entry at candle open" in r.flags


def test_stop_loss_exit_at_first_candle_low():
    r = E.run_day(D, day(body={(10, 0): (108, 111, 108, 110), (11, 0): (100, 100, 89, 95)}))
    assert r.exit_reason == E.STOP_LOSS and r.exit_price == 90 and r.exit_time == t(11, 0) and r.pnl_points == -20


def test_stop_gap_down_fills_at_open():
    r = E.run_day(D, day(body={(10, 0): (108, 111, 108, 110), (11, 0): (85, 86, 84, 85)}))
    assert r.exit_reason == E.STOP_LOSS and r.exit_price == 85


def test_export_shows_put_and_ambiguous_sides():
    from backend.strategies.breakout15.export import build_xlsx

    put = day(body={(10, 0): (92, 93, 88, 89), (15, 15): (80, 81, 79, 80)})
    d2 = D + dt.timedelta(days=1)
    both = [(x.replace(day=d2.day), *rest) for (x, *rest) in day(body={(10, 0): (100, 111, 89, 100)})]
    rows, summary = E.run_backtest(put + both)
    result = {"symbol": "BANKNIFTY", "token": "x", "from": "x", "to": "y", "generated_at": "x", "summary": summary,
              "rows": [r.as_dict() for r in rows]}
    ws = openpyxl.load_workbook(io.BytesIO(build_xlsx(result))).active
    assert ws["A1"].value == "BANK NIFTY"
    assert [ws["B4"].value, ws["B5"].value] == ["PUT", "UNRESOLVED"]
    assert ws["D4"].value == "PROFIT" and ws["E5"].value == "LOSS"


def test_candle_breaking_both_ends_without_minute_data_stays_unresolved():
    r = E.run_day(D, day(body={(10, 0): (100, 111, 89, 100)}))
    assert r.side == "BOTH" and r.exit_reason == E.STOP_LOSS and r.pnl_points == -20
    assert r.entry_price is None and any(f.startswith("AMBIGUOUS") for f in r.flags)
    assert round(r.pnl_pct, 4) == -20.0  # % of the first candle's open (100)


def minutes(start_hm, rows):
    """1-minute bars from `start_hm`: rows = [(o, h, l, c), ...] one per minute."""
    t0 = t(*start_hm)
    return [(t0 + dt.timedelta(minutes=i), *row) for i, row in enumerate(rows)]


BOTH_BODY = {(10, 0): (100, 111, 89, 100)}  # one 15-minute candle through both ends of 90..110


def test_both_ends_resolved_to_a_put_when_the_low_breaks_first_minute_by_minute():
    mins = minutes((10, 0), [(100, 101, 99, 100), (100, 100, 88, 89), (89, 100, 89, 99), (99, 112, 98, 111)])
    r = E.run_day(D, day(body=BOTH_BODY), minute_bars=mins)
    assert r.side == "PUT" and r.breakout_time == t(10, 1) and r.entry_price == 90
    assert r.exit_reason == E.STOP_LOSS and r.exit_time == t(10, 3) and r.exit_price == 110 and r.pnl_points == -20
    assert any(f.startswith("resolved on 1-minute") for f in r.flags) and not any(f.startswith("AMBIGUOUS") for f in r.flags)


def test_both_ends_resolved_to_a_call_when_the_high_breaks_first():
    mins = minutes((10, 0), [(100, 111, 100, 110), (110, 111, 89, 90)])
    r = E.run_day(D, day(body=BOTH_BODY), minute_bars=mins)
    assert r.side == "CALL" and r.entry_price == 110 and r.exit_price == 90 and r.pnl_points == -20
    assert r.breakout_time == t(10, 0) and r.exit_time == t(10, 1)


def test_one_minute_breaking_both_ends_picks_the_nearer_end_and_stays_flagged_ambiguous():
    mins = minutes((10, 0), [(105, 111, 89, 100)])  # opens nearer the High (110) than the Low (90)
    r = E.run_day(D, day(body=BOTH_BODY), minute_bars=mins)
    assert r.side == "CALL" and r.exit_reason == E.STOP_LOSS and any(f.startswith("AMBIGUOUS") for f in r.flags)


def test_minute_data_with_no_break_leaves_the_day_unresolved():
    r = E.run_day(D, day(body=BOTH_BODY), minute_bars=minutes((10, 0), [(100, 101, 99, 100)]))
    assert r.side == "BOTH"


def test_ambiguous_candles_lists_the_candle_to_fetch_and_run_backtest_uses_the_minutes():
    bars = day(body=BOTH_BODY)
    assert E.ambiguous_candles(bars) == {D: t(10, 0)}
    mins = {D: minutes((10, 0), [(100, 111, 100, 110), (110, 111, 89, 90)])}
    rows, summary = E.run_backtest(bars, minute_bars=mins)
    assert rows[0].side == "CALL" and summary["resolved_days"] == 1 and summary["call_entries"] == 1


def test_gap_open_above_high_then_falling_to_the_low_is_a_clear_call_stop_not_ambiguous():
    r = E.run_day(D, day(body={(9, 30): (115, 116, 85, 90)}))
    assert r.side == "CALL" and r.entry_price == 115 and r.exit_reason == E.STOP_LOSS
    assert r.exit_price == 90 and r.pnl_points == -25  # filled at the stop level (first low), never ambiguous
    assert not any(f.startswith("AMBIGUOUS") for f in r.flags)


def test_first_break_downwards_is_a_put_and_a_later_high_break_is_not_a_second_trade():
    r = E.run_day(D, day(body={(10, 0): (100, 100, 85, 90), (11, 0): (100, 112, 100, 111)}))
    assert r.side == "PUT" and r.breakout_time == t(10, 0) and r.entry_price == 90
    assert r.stop_loss == 110 and r.exit_reason == E.STOP_LOSS and r.exit_time == t(11, 0)
    assert r.exit_price == 110 and r.pnl_points == -20


def test_missing_first_candle_is_data_incomplete():
    r = E.run_day(D, day(skip={(9, 15)}))
    assert r.exit_reason == E.DATA_INCOMPLETE and r.entry_price is None and r.first_open is None


def test_missing_candle_hides_possible_breakout_flag_on_no_entry_day():
    r = E.run_day(D, day(skip={(12, 0)}))
    assert r.exit_reason == E.NO_ENTRY and any("cannot be ruled out" in f for f in r.flags)


def test_data_ending_early_is_not_labelled_market_close():
    r = E.run_day(D, day(body={(10, 0): (108, 111, 108, 110)}, until=(14, 0)))
    assert r.exit_reason == E.DATA_ENDED and any("not a verified exit" in f for f in r.flags)


def test_second_last_candle_percentage():
    r = E.run_day(D, day(body={(15, 0): (100, 103, 99, 102)}))
    assert r.second_last_time == t(15, 0) and r.second_last_close == 102 and round(r.second_last_change_pct, 2) == 2.0


def test_summary_counts_and_drawdown():
    days = [E.DayResult(D + dt.timedelta(days=i), side="CALL" if i % 2 == 0 else "PUT", entry_price=100, pnl_points=p,
                        exit_reason=E.MARKET_CLOSE if p > 0 else E.STOP_LOSS) for i, p in enumerate([10, -30, 5, -5])]
    days.append(E.DayResult(D + dt.timedelta(days=9)))
    s = E.summarize(days)
    assert s["trading_days"] == 5 and s["call_entries"] == 2 and s["put_entries"] == 2 and s["no_entry_days"] == 1
    assert s["winning_trades"] == 2 and s["losing_trades"] == 2 and s["win_rate"] == 50.0
    assert s["net_points"] == -20 and s["max_drawdown_points"] == 30 and s["market_close_exits"] == 2


def test_run_backtest_one_row_per_day():
    d2 = D + dt.timedelta(days=1)
    b2 = [(x.replace(day=d2.day), *rest) for (x, *rest) in day()]
    rows, s = E.run_backtest(day() + b2)
    assert [r.date for r in rows] == [D, d2] and s["trading_days"] == 2


def test_put_enters_at_first_low_and_profits_when_price_closes_lower():
    r = E.run_day(D, day(body={(10, 0): (92, 93, 88, 89), (15, 15): (80, 81, 79, 80)}))
    assert r.side == "PUT" and r.entry_price == 90 and r.stop_loss == 110
    assert r.exit_reason == E.MARKET_CLOSE and r.exit_price == 80 and r.pnl_points == 10
    assert round(r.pnl_pct, 4) == round(10 / 90 * 100, 4)


def test_put_gap_open_below_low_enters_at_open_and_stop_is_the_first_high():
    r = E.run_day(D, day(body={(9, 30): (85, 86, 84, 85), (11, 0): (100, 112, 99, 111)}))
    assert r.side == "PUT" and r.entry_price == 85 and "gap entry at candle open" in r.flags
    assert r.exit_reason == E.STOP_LOSS and r.exit_price == 110 and r.pnl_points == -25


def test_put_stop_gapping_above_fills_at_open():
    r = E.run_day(D, day(body={(10, 0): (92, 93, 88, 89), (11, 0): (115, 116, 114, 115)}))
    assert r.exit_reason == E.STOP_LOSS and r.exit_price == 115


def test_exit_mode_second_last_uses_the_1500_candle_close():
    body = {(10, 0): (108, 111, 108, 110), (15, 0): (112, 113, 111, 112), (15, 15): (120, 121, 119, 120)}
    close = E.run_day(D, day(body=body), "CLOSE")
    early = E.run_day(D, day(body=body), "SECOND_LAST")
    assert close.exit_reason == E.MARKET_CLOSE and close.exit_price == 120
    assert early.exit_reason == E.SECOND_LAST_CLOSE and early.exit_price == 112 and early.exit_time == t(15, 0)


def test_exit_mode_second_last_ignores_a_breakout_in_the_last_candle():
    r = E.run_day(D, day(body={(15, 15): (100, 115, 100, 112)}), "SECOND_LAST")
    assert r.exit_reason == E.NO_ENTRY


def test_invalid_exit_mode_is_rejected():
    import pytest

    with pytest.raises(ValueError):
        E.run_day(D, day(), "SOON")


# ---- data layer (backend/api/breakout15.py) ----

import asyncio

from backend.api import breakout15 as api


def test_chunks_cover_range_without_overlap():
    plan = api.chunk_ranges(dt.date(2025, 1, 1), dt.date(2025, 12, 31))
    assert plan[0][0].date() == dt.date(2025, 1, 1) and plan[-1][1].date() == dt.date(2025, 12, 31)
    for (_, b), (a, _) in zip(plan, plan[1:]):
        assert a.date() == b.date() + dt.timedelta(days=1)
    assert all((b - a).days < 200 for a, b in plan)


def test_to_bars_converts_ist_dedupes_sorts_and_clips():
    rows = [
        ["2025-01-07T09:30:00+05:30", 2, 3, 1, 2.5, 0],
        ["2025-01-07T09:15:00+05:30", 1, 2, 0.5, 1.5, 0],
        ["2025-01-07T09:15:00+05:30", 1, 2, 0.5, 1.5, 0],  # duplicate
        ["2025-02-01T09:15:00+05:30", 9, 9, 9, 9, 0],  # outside range
    ]
    bars = api.to_bars(rows, dt.date(2025, 1, 1), dt.date(2025, 1, 31))
    assert [b[0] for b in bars] == [dt.datetime(2025, 1, 7, 9, 15), dt.datetime(2025, 1, 7, 9, 30)]
    assert bars[0][1:] == (1.0, 2.0, 0.5, 1.5)


def test_fetch_bars_calls_each_chunk_once_and_reports_progress():
    calls, seen = [], []

    async def fake(a, b):
        calls.append((a, b))
        return [[f"{a:%Y-%m-%d}T09:15:00+05:30", 1, 2, 0.5, 1.5, 0]]

    bars = asyncio.run(api.fetch_bars("99926000", dt.date(2025, 1, 1), dt.date(2025, 8, 1), fake,
                                      progress=lambda d, t: seen.append((d, t)), delay=0))
    assert len(calls) == len(api.chunk_ranges(dt.date(2025, 1, 1), dt.date(2025, 8, 1))) == seen[-1][1]
    assert len(bars) == len(calls)


# ---- Excel import ----

import io

import openpyxl

from backend.strategies.breakout15.excel_import import ExcelImportError, parse_workbook


def _book(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEAD = ["DATE", "CALL/PUT", "15 minute candel size", "PROFIT", "LOSS", "EOD"]


def test_excel_parse_user_sheet_layout():
    content = _book([
        ["NIFTY 50", None, None, None, None, None], [None] * 6, HEAD,
        [dt.datetime(2025, 1, 1), "PUT ", 0.0031, None, "LOSS", 0.0031],
        [dt.datetime(2025, 1, 7), "CALL", 0.0038, None, "LOSS", 0.0038],
        [dt.datetime(2025, 1, 16), "put ", 0.0028, "PROFIT", None, 0.0005],
        ["junk", None, None, None, None, None],
    ])
    rows, skipped = parse_workbook(content)
    assert [r["date"] for r in rows] == ["2025-01-01", "2025-01-07", "2025-01-16"]
    assert [r["side"] for r in rows] == ["PUT", "CALL", "PUT"]
    assert rows[1]["candle_pct"] == 0.38 and rows[1]["outcome"] == "LOSS" and rows[2]["outcome"] == "PROFIT"
    assert rows[0]["eod_pct"] == 0.31 and len(skipped) == 1


def test_excel_percent_values_already_in_percent_are_kept():
    rows, _ = parse_workbook(_book([HEAD, [dt.datetime(2025, 1, 2), "CALL", 0.31, "PROFIT", None, 1.5]]))
    assert rows[0]["candle_pct"] == 0.31 and rows[0]["eod_pct"] == 1.5


def test_excel_without_date_header_or_bad_file_is_rejected():
    import pytest

    with pytest.raises(ExcelImportError):
        parse_workbook(_book([["a", "b"], [1, 2]]))
    with pytest.raises(ExcelImportError):
        parse_workbook(b"not an excel file")


def test_excel_one_outlier_does_not_flip_the_fraction_scaling():
    rows, _ = parse_workbook(_book([
        HEAD,
        [dt.datetime(2025, 1, 2), "PUT", 0.0025, None, "LOSS", 0.0025],
        [dt.datetime(2025, 1, 3), "PUT", 0.0028, "PROFIT", None, 0.005],
        [dt.datetime(2025, 1, 4), "PUT", 0.0023, None, "LOSS", 0.23],
    ]))
    assert [r["eod_pct"] for r in rows] == [0.25, 0.5, 23.0]


def test_export_xlsx_matches_the_users_manual_sheet_layout():
    from backend.strategies.breakout15.export import build_xlsx

    win = day(body={(10, 0): (108, 111, 108, 110), (15, 15): (115, 116, 114, 120)})
    d2 = D + dt.timedelta(days=1)
    loss = [(x.replace(day=d2.day), *rest) for (x, *rest) in day(body={(10, 0): (108, 111, 108, 110), (11, 0): (100, 100, 89, 95)})]
    d3 = D + dt.timedelta(days=2)
    flat = [(x.replace(day=d3.day), *rest) for (x, *rest) in day()]
    rows, summary = E.run_backtest(win + loss + flat)
    result = {"symbol": "NIFTY", "token": "99926000", "from": "x", "to": "y", "generated_at": "x", "summary": summary,
              "rows": [r.as_dict() for r in rows]}
    ws = openpyxl.load_workbook(io.BytesIO(build_xlsx(result))).active
    assert ws.title == "Sheet1" and ws["A1"].value == "NIFTY 50"
    assert [c.value for c in ws[3]] == ["DATE", "CALL/PUT", "15 minute candel size", "PROFIT", "LOSS", "EOD"]
    r4, r5, r6 = ([c.value for c in ws[i]] for i in (4, 5, 6))
    assert r4[0].date() == D and r4[1] == "CALL" and round(r4[2], 4) == 0.2 and r4[3] == "PROFIT" and r4[4] is None
    assert round(r4[5], 6) == round(10 / 110, 6)
    assert r5[1] == "CALL" and r5[3] is None and r5[4] == "LOSS" and round(r5[5], 6) == round(20 / 110, 6)
    assert r6[1] == "NO ENTRY" and r6[3] is None and r6[4] is None and r6[5] is None
    assert [c.value for c in ws[4]][1] == "CALL"
    assert ws["C4"].number_format == "0.00%" and ws["A4"].number_format == "d-mmm-yy" and ws.max_row == 6
