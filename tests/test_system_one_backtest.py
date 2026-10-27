"""System One trade backtest: engine rules, API response and Excel export.

All fixtures are fixed weekly closes — no Angel One calls. The expected
numbers were worked out by hand from the closes, not read back from the engine:

  FLAT      30 weeks oscillating around 100 (RSI ~50, 20 SMA = 100)
  RALLY     104, 108, ... 128 -> RSI first crosses above 60 on the 104 candle
            (index 30); entry is the NEXT candle, 108 (index 31).
  20 SMA at the entry candle = mean of closes[12:32] = 100.65
  risk per share = 108 - 100.65 = 7.35
  quantity = floor(10000 / 7.35) = floor(1360.54) = 1360
"""
from __future__ import annotations

import datetime as dt
import io
import math

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from backend.config.strategy_config import SystemOneBacktestConfig
from backend.indicators.moving_averages import sma
from backend.indicators.rsi import rsi
from backend.main import app
from backend.services import angelone_credential_store, system_one_backtest_service
from backend.services.top_bottom_backtest_service import ResolvedInstrument
from backend.strategies.system_one_backtest import SystemOneBacktestResult, SystemOneTrade, run_system_one_backtest

FLAT = [100, 101, 100, 99] * 7 + [100, 101]  # 30 weeks
RALLY = [104, 108, 112, 116, 120, 124, 126, 128]
CLIMB = [130 + 2 * k for k in range(14)]  # keeps the 20 SMA far below price for a long time

WINNER = FLAT + RALLY + CLIMB + [125, 118]  # exit candle: 125 (below its 20 SMA 137.05), index 52
LOSER = FLAT + RALLY + [110, 100]  # exit candle: 100 (below its 20 SMA 107.45), index 39
STILL_OPEN = FLAT + RALLY  # never closes below the 20 SMA
SIGNAL_ON_LAST_CANDLE = FLAT + [104]
ENTRY_BELOW_MA = FLAT + [104, 90]
NO_SIGNAL = FLAT
# RSI crosses 60 at index 16, where the 20 SMA does not exist yet (first value at index 19)
MA_NOT_READY = [100 - i for i in range(15)] + [96, 106, 116, 126, 136, 146]
# second fresh cross (index 41) while the first trade is still open
SECOND_CROSS_WHILE_OPEN = FLAT + RALLY + [124, 118, 114, 120, 128, 136, 140]

SIGNAL_IDX, ENTRY_IDX = 30, 31


def weekly(closes: list[float], scale: float = 1.0) -> pd.DataFrame:
    idx = pd.date_range("2024-01-05", periods=len(closes), freq="W-FRI")
    c = pd.Series([x * scale for x in closes], index=idx, dtype=float)
    return pd.DataFrame({"open": c, "high": c, "low": c, "close": c, "volume": 1000.0})


def hourly(closes: list[float]) -> pd.DataFrame:
    idx = pd.date_range("2024-01-05 09:15", periods=len(closes), freq="60min")
    c = pd.Series(closes, index=idx, dtype=float)
    return pd.DataFrame({"open": c, "high": c, "low": c, "close": c, "volume": 1000.0})


def cross_indexes(df: pd.DataFrame, level: float = 60.0) -> list[int]:
    r = rsi(df["close"], 14)
    return [i for i in range(1, len(df)) if r.iloc[i - 1] <= level < r.iloc[i]]


# ---- signal -----------------------------------------------------------------

def test_fixture_crosses_60_only_once_at_signal_candle():
    df = weekly(WINNER)
    assert cross_indexes(df) == [SIGNAL_IDX]


def test_buy_signal_is_a_fresh_cross_not_just_rsi_above_60():
    # RSI stays above 60 for many candles after the cross (31..51) — only ONE trade.
    result = run_system_one_backtest(weekly(WINNER))
    assert len(result.trades) == 1
    assert result.trades[0].signal_date == weekly(WINNER).index[SIGNAL_IDX]


def test_no_signal_gives_no_trades_and_no_warning():
    result = run_system_one_backtest(weekly(NO_SIGNAL))
    assert result.trades == []
    assert result.skipped == []
    assert result.warnings == []


def test_second_cross_while_trade_open_is_ignored():
    df = weekly(SECOND_CROSS_WHILE_OPEN)
    assert len(cross_indexes(df)) == 2  # the fixture really has two fresh crosses
    result = run_system_one_backtest(df)
    assert len(result.trades) == 1
    assert result.trades[0].status == "OPEN"


def test_engine_gives_identical_trade_numbers_on_any_candle_series():
    # Same closes, hourly stamps instead of weekly: the rules only look at the
    # candle sequence, so entry/stop/quantity/P&L must not change.
    t = run_system_one_backtest(hourly(WINNER)).trades[0]
    assert (t.entry_price, t.quantity, t.exit_price) == (108.0, 1360, 125.0)
    assert t.stop_loss == pytest.approx(100.65)
    assert t.pnl_amount == pytest.approx(23_120)
    assert t.entry_date == hourly(WINNER).index[ENTRY_IDX]


# ---- entry ------------------------------------------------------------------

def test_entry_is_next_candles_close_and_date():
    df = weekly(WINNER)
    t = run_system_one_backtest(df).trades[0]
    assert t.signal_date == df.index[SIGNAL_IDX]
    assert t.entry_date == df.index[ENTRY_IDX]
    assert t.entry_price == 108.0  # next candle's close ...
    assert t.entry_price != df["close"].iloc[SIGNAL_IDX]  # ... not the signal candle's 104


def test_signal_on_last_candle_creates_no_trade():
    result = run_system_one_backtest(weekly(SIGNAL_ON_LAST_CANDLE))
    assert result.trades == []
    assert len(result.skipped) == 1
    assert "next candle" in result.skipped[0].reason


# ---- stop loss, risk, quantity -----------------------------------------------

def test_stop_loss_is_20_sma_at_entry_and_not_overwritten_later():
    df = weekly(WINNER)
    t = run_system_one_backtest(df).trades[0]
    expected = df["close"].iloc[ENTRY_IDX - 19 : ENTRY_IDX + 1].mean()
    assert t.stop_loss == pytest.approx(expected)
    assert t.stop_loss == pytest.approx(100.65)
    # the 20 SMA kept rising until the exit candle; the saved stop did not follow it
    ma_at_exit = sma(df["close"], 20).iloc[52]
    assert ma_at_exit == pytest.approx(137.05)
    assert t.stop_loss != pytest.approx(ma_at_exit)


def test_risk_per_share_risk_per_trade_and_quantity_rounded_down():
    t = run_system_one_backtest(weekly(WINNER)).trades[0]
    assert t.risk_per_share == pytest.approx(7.35)
    assert t.risk_per_trade == 10_000.0
    assert t.quantity == 1360  # floor(1360.54), not rounded up to 1361
    assert t.quantity * t.risk_per_share <= 10_000.0 < (t.quantity + 1) * t.risk_per_share


def test_default_config_risks_one_percent_of_ten_lakh():
    cfg = SystemOneBacktestConfig()
    assert cfg.reference_capital == 1_000_000.0
    assert cfg.risk_per_trade_pct == 1.0
    assert cfg.risk_per_trade == 10_000.0
    assert math.floor(10_000 / 167) == 59  # the worked example from the spec


def test_quantity_follows_configured_capital():
    cfg = SystemOneBacktestConfig(reference_capital=100_000.0)  # risk per trade 1,000
    t = run_system_one_backtest(weekly(WINNER), cfg).trades[0]
    assert t.risk_per_trade == 1_000.0
    assert t.quantity == 136  # floor(1000 / 7.35)


def test_entry_not_above_20_sma_is_skipped_not_traded():
    result = run_system_one_backtest(weekly(ENTRY_BELOW_MA))
    assert result.trades == []
    assert len(result.skipped) == 1
    assert "not above the 20 MA" in result.skipped[0].reason


def test_zero_quantity_is_skipped_with_reason():
    # same shape, prices x2000: risk per share 14,700 > risk per trade 10,000
    result = run_system_one_backtest(weekly(WINNER, scale=2000))
    assert result.trades == []
    assert len(result.skipped) == 1
    assert "Quantity would be 0" in result.skipped[0].reason


def test_missing_20_ma_at_entry_is_skipped_without_inventing_a_stop():
    df = weekly(MA_NOT_READY)
    assert cross_indexes(df) == [16]
    result = run_system_one_backtest(df)
    assert result.trades == []
    assert "20 MA not available" in result.skipped[0].reason


def test_insufficient_history_returns_warning_not_crash():
    result = run_system_one_backtest(weekly([100, 101, 102, 103, 104, 105, 106, 107, 108, 109]))
    assert result.trades == []
    assert "Insufficient history" in result.warnings[0]


def test_empty_or_missing_data_returns_warning_not_crash():
    assert run_system_one_backtest(pd.DataFrame()).warnings
    assert run_system_one_backtest(None).warnings


# ---- exit and P&L -------------------------------------------------------------

def test_exit_on_first_close_below_20_sma_profit():
    df = weekly(WINNER)
    t = run_system_one_backtest(df).trades[0]
    assert t.status == "CLOSED"
    assert t.exit_date == df.index[52]
    assert t.exit_price == 125.0
    assert t.pnl_amount == pytest.approx((125.0 - 108.0) * 1360)  # +23,120
    assert t.pnl_amount > 0


def test_exit_loss_is_negative():
    df = weekly(LOSER)
    t = run_system_one_backtest(df).trades[0]
    assert t.status == "CLOSED"
    assert t.exit_date == df.index[39]
    assert t.exit_price == 100.0
    assert t.pnl_amount == pytest.approx((100.0 - 108.0) * 1360)  # -10,880
    assert t.pnl_amount < 0


def test_close_above_20_sma_does_not_exit():
    # index 38 closes at 110, still above its 20 SMA (107.40): no exit there
    t = run_system_one_backtest(weekly(LOSER)).trades[0]
    assert t.exit_date != weekly(LOSER).index[38]


def test_trade_with_no_exit_stays_open_without_fabricated_exit_or_pnl():
    t = run_system_one_backtest(weekly(STILL_OPEN)).trades[0]
    assert t.status == "OPEN"
    assert t.exit_price is None and t.exit_date is None and t.pnl_amount is None
    s = run_system_one_backtest(weekly(STILL_OPEN)).summary()
    assert (s.open_trades, s.closed_trades, s.net_realized_pnl) == (1, 0, 0.0)


def test_summary_counts_only_closed_trades():
    s = run_system_one_backtest(weekly(LOSER)).summary()
    assert (s.total_trades, s.closed_trades, s.winning_trades, s.losing_trades) == (1, 1, 0, 1)
    assert s.net_realized_pnl == pytest.approx(-10_880)


# ---- date window ---------------------------------------------------------------

def test_signal_before_from_date_is_ignored():
    df = weekly(WINNER)
    result = run_system_one_backtest(df, from_date=df.index[SIGNAL_IDX] + dt.timedelta(days=1))
    assert result.trades == []


def test_candles_after_to_date_are_dropped_so_trade_stays_open():
    df = weekly(WINNER)
    result = run_system_one_backtest(df, to_date=df.index[40].to_pydatetime())
    assert result.trades[0].status == "OPEN"


# ---- chart series and summary statistics --------------------------------------

def test_series_covers_every_candle_with_indicator_values_and_none_during_warmup():
    df = weekly(WINNER)
    series = run_system_one_backtest(df).series
    assert len(series) == len(df)
    assert series[0].rsi is None and series[0].ma is None  # not enough history yet
    s31 = series[ENTRY_IDX]
    assert s31.close == 108.0
    assert s31.ma == pytest.approx(100.65)  # the same 20 SMA the stop loss was taken from
    assert s31.rsi == pytest.approx(rsi(df["close"], 14).iloc[ENTRY_IDX])


def test_series_is_limited_to_the_requested_window():
    df = weekly(WINNER)
    from_date = df.index[20].to_pydatetime()
    to_date = df.index[45].to_pydatetime()
    series = run_system_one_backtest(df, from_date=from_date, to_date=to_date).series
    assert series[0].date == df.index[20]
    assert series[-1].date == df.index[45]
    assert len(series) == 26
    assert series[0].ma is not None  # warm-up candles before the window still feed the 20 SMA


def _closed(n: int, pnl: float) -> SystemOneTrade:
    ts = pd.Timestamp("2024-01-05")
    return SystemOneTrade(
        trade_number=n, signal_date=ts, entry_date=ts, entry_price=100.0, stop_loss=90.0, risk_per_share=10.0,
        risk_per_trade=10_000.0, quantity=1, status="CLOSED", exit_date=ts, exit_price=100.0, pnl_amount=pnl,
    )


def test_summary_statistics_from_mixed_trades():
    open_trade = _closed(4, 0.0)
    open_trade.status, open_trade.exit_price, open_trade.exit_date, open_trade.pnl_amount = "OPEN", None, None, None
    s = SystemOneBacktestResult(trades=[_closed(1, 300.0), _closed(2, -100.0), _closed(3, 100.0), open_trade]).summary()
    assert (s.total_trades, s.closed_trades, s.open_trades) == (4, 3, 1)
    assert s.win_rate == pytest.approx(200 / 3)
    assert s.average_win == pytest.approx(200.0)
    assert s.average_loss == pytest.approx(-100.0)
    assert s.profit_factor == pytest.approx(4.0)  # 400 gross profit / 100 gross loss
    assert (s.best_trade, s.worst_trade) == (300.0, -100.0)
    assert s.net_realized_pnl == pytest.approx(300.0)


def test_summary_statistics_are_none_not_zero_when_there_is_nothing_to_compute():
    s = run_system_one_backtest(weekly(NO_SIGNAL)).summary()
    assert s.win_rate is None and s.profit_factor is None and s.best_trade is None
    only_win = run_system_one_backtest(weekly(WINNER)).summary()
    assert only_win.win_rate == 100.0 and only_win.average_loss is None and only_win.profit_factor is None
    only_open = run_system_one_backtest(weekly(STILL_OPEN)).summary()
    assert only_open.win_rate is None


# ---- API + Excel -----------------------------------------------------------------

@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(angelone_credential_store, "CREDENTIAL_PATH", tmp_path / "angelone_credentials.json")
    with TestClient(app) as c:
        yield c


def _patch_data(monkeypatch, closes, frame=weekly, seen=None):
    async def fake_resolve(symbol, instrument_type, expiry):
        return ResolvedInstrument(
            underlying=symbol, trading_symbol=f"{symbol}-EQ", token="1", exch_seg="NSE",
            expiry="", instrument_class="STOCK", lot_size=1,
        )

    async def fake_fetch(angelone, instrument, timeframe, from_date=None):
        if seen is not None:
            seen.append(timeframe)
        return frame(closes)

    monkeypatch.setattr(system_one_backtest_service, "_resolve_instrument", fake_resolve)
    monkeypatch.setattr(system_one_backtest_service, "_fetch_ohlcv", fake_fetch)


_BODY = {"symbol": "TESTCO", "from_date": "2024-01-01T00:00:00", "to_date": "2026-12-31T23:59:59"}


def test_api_returns_real_trade_rows(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    resp = client.post("/api/system-one/backtest", json=_BODY)
    assert resp.status_code == 200
    body = resp.json()
    assert body["timeframe"] == "1W" and body["risk_per_trade"] == 10_000.0
    assert len(body["trades"]) == 1
    t = body["trades"][0]
    assert t["entry_date"].startswith("2024-08-09")  # weekly index[31]: 2024-01-05 + 31 weeks
    assert t["entry_price"] == 108.0
    assert t["stop_loss"] == pytest.approx(100.65)
    assert t["risk_per_share"] == pytest.approx(7.35)
    assert t["quantity"] == 1360
    assert t["exit_price"] == 125.0
    assert t["pnl_amount"] == pytest.approx(23_120)
    assert body["summary"]["net_realized_pnl"] == pytest.approx(23_120)
    assert "gross" in body["note"]


def test_api_returns_chart_candles_and_summary_stats(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    body = client.post("/api/system-one/backtest", json=_BODY).json()
    assert len(body["candles"]) == len(WINNER)
    first, entry = body["candles"][0], body["candles"][ENTRY_IDX]
    assert first["rsi"] is None and first["ma"] is None  # JSON null, not NaN
    assert entry["close"] == 108.0 and entry["ma"] == pytest.approx(100.65)
    assert body["summary"]["win_rate"] == 100.0
    assert body["summary"]["profit_factor"] is None


def test_api_timeframe_defaults_to_weekly(client, monkeypatch):
    seen: list[str] = []
    _patch_data(monkeypatch, WINNER, seen=seen)
    body = client.post("/api/system-one/backtest", json=_BODY).json()
    assert seen == ["1W"] and body["timeframe"] == "1W"


@pytest.mark.parametrize("tf", ["1H", "4H", "1D", "1W"])
def test_api_passes_the_chosen_timeframe_to_the_data_fetch(client, monkeypatch, tf):
    seen: list[str] = []
    _patch_data(monkeypatch, WINNER, seen=seen)
    body = client.post("/api/system-one/backtest", json={**_BODY, "timeframe": tf}).json()
    assert seen == [tf] and body["timeframe"] == tf
    assert body["trades"][0]["quantity"] == 1360


@pytest.mark.parametrize("tf", ["1m", "15m", "1M", "2W", ""])
def test_api_rejects_unsupported_timeframes_with_400(client, monkeypatch, tf):
    _patch_data(monkeypatch, WINNER)
    assert client.post("/api/system-one/backtest", json={**_BODY, "timeframe": tf}).status_code == 400


def test_api_warns_when_data_starts_after_the_requested_from_date(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    body = client.post("/api/system-one/backtest", json={**_BODY, "from_date": "2023-01-01T00:00:00"}).json()
    assert any("only available from 2024-01-05" in w for w in body["warnings"])
    ok = client.post("/api/system-one/backtest", json=_BODY).json()  # starts 2024-01-05, inside the slack
    assert ok["warnings"] == []


def test_api_open_trade_has_null_exit_fields(client, monkeypatch):
    _patch_data(monkeypatch, STILL_OPEN)
    t = client.post("/api/system-one/backtest", json=_BODY).json()["trades"][0]
    assert t["status"] == "OPEN"
    assert t["exit_price"] is None and t["exit_date"] is None and t["pnl_amount"] is None


def test_api_no_trades_is_a_normal_empty_result(client, monkeypatch):
    _patch_data(monkeypatch, NO_SIGNAL)
    resp = client.post("/api/system-one/backtest", json=_BODY)
    assert resp.status_code == 200
    assert resp.json()["trades"] == []


def test_api_rejects_reversed_dates_with_400(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    resp = client.post("/api/system-one/backtest", json={**_BODY, "from_date": "2027-01-01T00:00:00"})
    assert resp.status_code == 400


def _export_rows(client, backtest_id):
    resp = client.get(f"/api/system-one/backtest/{backtest_id}/export")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Trades"]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_excel_export_matches_ui_columns_and_values(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    bid = client.post("/api/system-one/backtest", json=_BODY).json()["backtest_id"]
    rows = _export_rows(client, bid)
    assert rows[0][:5] == ["Trade #", "Signal Date", "Entry Date", "Entry Price", "Initial Stop (MA at entry)"]
    r = rows[1]
    assert r[0] == 1
    assert r[2].date() == weekly(WINNER).index[31].date()
    assert r[3] == 108.0 and r[6] == 1360 and r[9] == 125.0
    assert r[4] == pytest.approx(100.65) and r[5] == pytest.approx(7.35) and r[7] == 10_000.0
    assert r[8].date() == weekly(WINNER).index[52].date()
    assert r[10] == "GAP_OPEN_BELOW_MA" and r[11] == 21
    assert r[12] == pytest.approx(23_120)
    assert r[13] == pytest.approx((125 - 108) / 108 * 100) and r[14] == pytest.approx(2.312)
    assert len(rows) == 2


def test_excel_export_keeps_time_of_day_for_intraday_and_plain_dates_otherwise(client, monkeypatch):
    _patch_data(monkeypatch, WINNER, frame=hourly)
    bid = client.post("/api/system-one/backtest", json={**_BODY, "timeframe": "1H"}).json()["backtest_id"]
    r = _export_rows(client, bid)[1]
    assert r[2] == hourly(WINNER).index[ENTRY_IDX].to_pydatetime()
    assert r[2].hour != 0 or r[2].minute != 0  # a real time of day, not midnight
    assert r[8] == hourly(WINNER).index[52].to_pydatetime()

    _patch_data(monkeypatch, WINNER, frame=weekly)
    bid = client.post("/api/system-one/backtest", json={**_BODY, "timeframe": "1W"}).json()["backtest_id"]
    assert _export_rows(client, bid)[1][2].hour == 0


def test_excel_export_keeps_negative_pnl(client, monkeypatch):
    _patch_data(monkeypatch, LOSER)
    bid = client.post("/api/system-one/backtest", json=_BODY).json()["backtest_id"]
    assert _export_rows(client, bid)[1][12] == pytest.approx(-10_880)


def test_excel_export_open_trade_leaves_exit_cells_empty(client, monkeypatch):
    _patch_data(monkeypatch, STILL_OPEN)
    bid = client.post("/api/system-one/backtest", json=_BODY).json()["backtest_id"]
    r = _export_rows(client, bid)[1]
    assert r[8] is None and r[9] is None and r[12] is None
    assert r[10] == "OPEN"
    assert r[15] is not None  # unrealised P&L sits in its own column


def test_excel_export_with_no_trades_is_header_only(client, monkeypatch):
    _patch_data(monkeypatch, NO_SIGNAL)
    bid = client.post("/api/system-one/backtest", json=_BODY).json()["backtest_id"]
    assert len(_export_rows(client, bid)) == 1


def test_export_unknown_id_is_404(client):
    assert client.get("/api/system-one/backtest/does-not-exist/export").status_code == 404


# ---- spec 2026-10: exit at the PREVIOUS candle's 20 SMA ------------------------
def ohlc(closes, opens=None, lows=None, highs=None) -> pd.DataFrame:
    """Weekly candles where open/low/high default to the close; override single candles by index."""
    df = weekly(closes)
    for col, over in (("open", opens), ("low", lows), ("high", highs)):
        for k, v in (over or {}).items():
            df.iloc[k, df.columns.get_loc(col)] = float(v)
    return df


def test_exit_when_only_low_touches_the_stop_exits_at_the_stop_level():
    # candle 38 closes 110 (above MA) but its LOW dips under the stop = MA of candle 37
    df = ohlc(LOSER, lows={38: 90})
    ma37 = sma(df["close"], 20).iloc[37]
    t = run_system_one_backtest(df).trades[0]
    assert t.exit_date == df.index[38]
    assert t.exit_price == pytest.approx(ma37)  # the stop level, not the 90 low, not the 110 close
    assert t.exit_reason == "STOP_AT_MA"
    assert t.pnl_amount == pytest.approx((ma37 - 108) * 1360)
    assert "assumed" in t.data_flag


def test_low_above_the_stop_does_not_exit():
    df = ohlc(LOSER, lows={38: 108})  # MA of candle 37 (106.9) is lower than 108
    assert sma(df["close"], 20).iloc[37] < 108
    assert run_system_one_backtest(df).trades[0].exit_date == df.index[39]


def test_gap_down_exits_at_the_open_not_at_the_stop():
    df = ohlc(LOSER, opens={38: 90}, lows={38: 85})
    ma37 = sma(df["close"], 20).iloc[37]
    t = run_system_one_backtest(df).trades[0]
    assert t.exit_date == df.index[38]
    assert t.exit_price == 90.0 and t.exit_price < ma37
    assert t.exit_reason == "GAP_OPEN_BELOW_MA"


def test_no_exit_on_the_entry_candle_even_if_its_low_is_below_the_ma():
    df = ohlc(STILL_OPEN, lows={ENTRY_IDX: 50})
    t = run_system_one_backtest(df).trades[0]
    assert t.entry_date == df.index[ENTRY_IDX] and t.status == "OPEN"


def test_stop_uses_previous_candle_ma_not_the_same_candles_ma():
    # candle 39: open 109 (above MA38), low 107.42 sits above MA38 (107.40) but below MA39 - must NOT exit
    df = ohlc(FLAT + RALLY + [110, 108], opens={39: 109}, lows={39: 107.42})
    ma = sma(df["close"], 20)
    assert ma.iloc[38] < 107.42
    assert run_system_one_backtest(df).trades[0].exit_date != df.index[39]


def test_pnl_pct_r_multiple_and_holding_periods():
    t = run_system_one_backtest(weekly(WINNER)).trades[0]
    assert t.holding_periods == 52 - ENTRY_IDX
    assert t.pnl_pct == pytest.approx(23_120 / (108 * 1360) * 100)
    assert t.r_multiple == pytest.approx(2.312)


def test_gap_loss_can_exceed_planned_risk_and_is_reported_in_full():
    t = run_system_one_backtest(ohlc(FLAT + RALLY[:2] + [70], opens={32: 70}, lows={32: 65})).trades[0]
    assert t.exit_price == 70.0
    assert t.pnl_amount == pytest.approx(-51_680)
    assert t.r_multiple == pytest.approx(-5.168)


def test_open_trade_unrealised_pnl_is_separate_and_not_in_realised_total():
    r = run_system_one_backtest(weekly(STILL_OPEN))
    t = r.trades[0]
    assert t.unrealized_pnl == pytest.approx((128 - 108) * 1360)
    s = r.summary()
    assert s.net_realized_pnl == 0.0 and s.unrealized_pnl == pytest.approx(27_200)


def test_one_trade_at_a_time_then_next_signal_after_exit():
    closes = FLAT + RALLY + [110, 100] + [96, 90, 85, 80, 78, 76, 74, 72, 70, 68, 66, 64, 66, 70, 80, 95, 110, 125, 140]
    r = run_system_one_backtest(weekly(closes))
    assert len(r.trades) >= 2
    assert r.trades[1].signal_date > r.trades[0].exit_date
    assert r.trades[1].entry_date > r.trades[0].exit_date


def test_max_drawdown_losing_streak_and_year_table():
    def closed(n, pnl, exit_date):
        t = _closed(n, pnl)
        t.exit_date = pd.Timestamp(exit_date)
        return t

    res = SystemOneBacktestResult(trades=[
        closed(1, 500.0, "2024-03-01"), closed(2, -200.0, "2024-06-01"), closed(3, -300.0, "2024-09-01"),
        closed(4, 100.0, "2025-02-01"),
    ])
    s = res.summary()
    assert s.max_drawdown == pytest.approx(500.0)  # peak 500 -> trough 0
    assert s.longest_losing_streak == 2
    y = {r.year: r for r in res.yearly()}
    assert (y[2024].trades, y[2024].wins, y[2024].losses, y[2024].net_pnl) == (3, 1, 2, 0.0)
    assert y[2024].gross_profit == 500.0 and y[2024].gross_loss == -500.0
    assert y[2025].net_pnl == 100.0
    assert [p.cumulative_pnl for p in res.equity_curve()] == [500.0, 300.0, 0.0, 100.0]


def test_data_gap_is_detected_and_flags_the_trade_it_spans():
    from backend.strategies.system_one_backtest import find_data_gaps, flag_trades_in_gaps

    df = weekly(WINNER)
    holed = df.drop(df.index[40:42])  # two missing weeks inside the trade
    gaps = find_data_gaps(holed.index, max_days=10)
    assert len(gaps) == 1 and gaps[0] == (df.index[39], df.index[42])
    r = run_system_one_backtest(holed)
    flag_trades_in_gaps(r, gaps)
    assert r.trades[0].data_flag.startswith("DATA GAP")
    assert find_data_gaps(df.index, max_days=10) == []


def test_api_exposes_new_fields_yearly_equity_curve_and_dated_file_name(client, monkeypatch):
    _patch_data(monkeypatch, WINNER)
    body = client.post("/api/system-one/backtest", json=_BODY).json()
    t = body["trades"][0]
    assert t["exit_reason"] == "GAP_OPEN_BELOW_MA" and t["holding_periods"] == 21
    assert t["r_multiple"] == pytest.approx(2.312) and t["unrealized_pnl"] is None
    assert body["yearly"][0]["year"] == 2025 and body["yearly"][0]["net_pnl"] == pytest.approx(23_120)
    assert body["equity_curve"][-1]["cumulative_pnl"] == pytest.approx(body["summary"]["net_realized_pnl"])
    assert body["summary"]["max_drawdown"] == 0.0
    exp = client.get(f"/api/system-one/backtest/{body['backtest_id']}/export")
    assert "system_one_TESTCO_1W_20240101_20261231.xlsx" in exp.headers["content-disposition"]


def test_excel_total_matches_ui_total(client, monkeypatch):
    _patch_data(monkeypatch, LOSER)
    body = client.post("/api/system-one/backtest", json=_BODY).json()
    resp = client.get(f"/api/system-one/backtest/{body['backtest_id']}/export")
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Summary"]
    vals = {r[0].value: r[1].value for r in ws.iter_rows()}
    assert vals["Net Realised P&L (closed trades only)"] == pytest.approx(body["summary"]["net_realized_pnl"])
